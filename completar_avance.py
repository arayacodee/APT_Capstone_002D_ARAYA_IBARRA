from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from pathlib import Path

out=Path(r'C:\Users\marti\Documents\Cronos\entregables')
out.mkdir(exist_ok=True)
d=Document(r'C:\Users\marti\Downloads\Ficha_Avance_04_Cronos_BORRADOR.docx')
p=d.paragraphs
repl={
1:'Proyecto Capstone CRONOS  Trazabilidad y planificación de la versión 1.0',
3:'CRONOS permite a clientes reservar servicios y a profesionales gestionar su negocio, disponibilidad y atenciones. Esta ficha relaciona el desarrollo existente con el cierre de la versión 1.0 y organiza las ocho semanas restantes.',
5:'Estimación inicial de desarrollo para planificación:',
6:'80 % como estimación propuesta del equipo antes de la validación funcional.',
8:'60 % de avance reconocido al cerrar la iteración 04, con evidencias de los flujos de acceso y sesión.',
10:'Validar de extremo a extremo la integración Angular–API–Neon y los servicios Google y Cloudinary. El desarrollo existe, pero faltan evidencias reproducibles de los flujos completos en un ambiente de desarrollo.',
12:'La revisión del repositorio identifica los módulos de la versión 1.0 y pruebas automatizadas. Se reconoce un cumplimiento parcial del 50 % por módulo mientras falta la demostración integrada. Este porcentaje mide el cierre comprobado del alcance, no la cantidad de código escrito.',
13:'Avance reconocido = peso × cumplimiento. Los pesos suman 100 %. Corte de revisión: 08-10-2026. Las pruebas aisladas respaldan reglas específicas; no sustituyen la validación con servicios reales.',
16:'Diferencia: 80 % estimado − 50 % reconocido = 30 puntos porcentuales pendientes de acreditar o completar.',
17:'¿El alcance puede completarse en ocho semanas? Sí con ajustes: mantener el MVP de reservas e historial y cerrar la validación antes de ampliar funciones.',
19:'Cada semana cerrará con una demostración y un registro de pruebas. Integrante 1 liderará frontend y coordinación; Manuel liderará backend e integraciones. Ambos revisarán los cambios y validarán el flujo completo.',
24:'Avance reconocido al corte de la ficha: 50 %. Meta de la iteración 04: 60 %.',
25:'Estado general de la planificación: Viable con ajustes.',
26:'Validación: Representante del equipo ____________________    Docente guía ____________________',
}
for i,t in repl.items(): p[i].text=t
d.tables[0].cell(0,1).text='Integrante 1: coordinación, frontend y validación funcional. Manuel: backend, datos e integraciones. Ambos: revisión y pruebas.'
d.tables[0].cell(0,3).text=''
rows=[
['Autenticación Google','Acceso de usuarios',10,'auth.controller.ts; login/login.ts. Pendiente demo con Google.'],
['Sesión y permisos','Acceso por rol',10,'auth.middleware.ts; lib/session.ts; guardias Angular. Pendiente comprobar sesión y cierre de sesión.'],
['API y gestión de servicios y disponibilidad','Operación del profesional',10,'service.controller.ts; availability.controller.ts; pantallas de servicios y disponibilidad. Pendiente flujo integrado.'],
['Interfaz Angular y navegación','Uso de cliente y profesional',10,'frontend/src/app/app.routes.ts y pages/. Pendiente validación de navegación, errores y vista móvil.'],
['Persistencia PostgreSQL Neon y Prisma','Datos del sistema',10,'backend/prisma/schema.prisma y migrations/. Pendiente validar migraciones y persistencia en Neon de desarrollo.'],
['Carga y eliminación de imágenes','Perfil y evidencias',5,'lib/cloudinary-upload.ts; business.controller.ts; session.controller.ts. Pendiente carga real y manejo de fallos.'],
['Gestión de negocio y perfil','Configuración profesional',10,'business.controller.ts; professional-onboarding/ y professional-profile/. Pendiente demo de creación y edición.'],
['Reservas y estados de atención','Flujo principal',25,'reservation.controller.ts; booking/ y my-reservations/. Pruebas de estados aprobadas; pendiente concurrencia real.'],
['Historial privado de atenciones','Seguimiento de clientes',10,'session.controller.ts; professional-history/. Pruebas de permisos aprobadas; pendiente demo de notas e imágenes.'],
]
for r,x in zip(d.tables[2].rows[1:],rows):
    name,origin,w,e=x
    for c,t in zip(r.cells,[name,origin,f'{w} %','Parcial\n50 %',f'{w/2:g} %',e]):c.text=t
summary=[['Funcionalidades terminadas y comprobadas','0','0 %','Sin validación integrada completa al corte.'],['Funcionalidades parciales o iniciadas','9','100 %','9 parciales al 50 %; ninguna solo iniciada.'],['Funcionalidades no iniciadas','0','0 %','Todos los módulos tienen implementación.'],['Avance total reconocido','9 módulos','50 % de 100 %','100 % × 50 % = 50 %.']]
for r,vals in zip(d.tables[3].rows[1:],summary):
    for c,t in zip(r.cells,vals):c.text=t
plan=[
['1\n09–15 oct.','Acordar alcance y demostrar acceso y sesión.','Integrante 1 y Manuel','Login Google, sesión, roles y logout con evidencia. Pesos aprobados.','60 %'],
['2\n16–22 oct.','Validar persistencia e imágenes.','Manuel; apoyo Integrante 1','Migraciones en desarrollo; datos persistentes; subida y eliminación real de imágenes.','67,5 %'],
['3\n23–29 oct.','Cerrar gestión de servicios, disponibilidad, negocio y perfil.','Manuel: API. Integrante 1: UI.','Crear/editar servicio y perfil; publicar horarios; controlar permisos y errores.','77,5 %'],
['4\n30 oct.–05 nov.','Validar el flujo principal de reservas.','Ambos','Reservar horario, evitar solapamientos, cancelar y completar una única atención.','90 %'],
['5\n06–12 nov.','Integrar y validar historial privado.','Integrante 1: UI. Manuel: API.','Atención en historial; notas, imágenes y paginación; acceso solo del propietario.','95 %'],
['6\n13–19 nov.','Cerrar validación transversal de la interfaz y documentar.','Integrante 1; revisión Manuel','Flujos cliente/profesional, vista móvil y errores verificados; manual y casos de prueba.','100 %'],
['7\n20–26 nov.','Corregir errores y ejecutar regresión.','Ambos','Versión candidata sin errores que impidan el flujo principal; repetir casos críticos.','100 %'],
['8\n27 nov.–03 dic.','Validación final y entrega de la versión 1.0.','Integrante 1: entrega. Manuel: entorno.','Demo completa, repositorio y guía de instalación; conformidad con criterios de aceptación.','100 %'],
]
for r,vals in zip(d.tables[4].rows[1:],plan):
    for c,t in zip(r.cells,vals):c.text=t
viability=[
['Alcance definido para la versión 1.0','MVP: acceso, negocio, servicios, horarios, reservas e historial privado. Sin administrador global.','Viable'],
['Capacidad del equipo para ocho semanas','Responsables complementarios; revisión semanal de pendientes y disponibilidad.','Ajustar carga semanal'],
['Dependencias técnicas o externas','Google, Neon y Cloudinary requieren pruebas reales y configuración del ambiente de desarrollo.','En riesgo hasta validación'],
['Tiempo reservado para pruebas','Semanas 5 y 6 para integración y pruebas; semana 7 para regresión; semana 8 para aceptación.','Suficiente con seguimiento'],
['Funciones que deben postergarse','Mantener fuera del MVP nuevas funciones como pagos, notificaciones y administración global.','Identificadas'],
]
for r,vals in zip(d.tables[6].rows[1:],viability):
    for c,t in zip(r.cells,vals):c.text=t
risks=[
['Integración Google, Neon y Cloudinary sin demo completa','Puede bloquear acceso, persistencia o imágenes.','Validar servicios reales en desarrollo; registrar resultados sin exponer credenciales.','Manuel y apoyo Integrante 1\n22-10-2026'],
['Concurrencia y estados de reservas','Solapamientos o atención duplicada afectarían el flujo principal.','Probar solicitudes simultáneas en PostgreSQL y completar/cancelar repetidamente.','Manuel; validación Integrante 1\n05-11-2026'],
['Conflictos de ramas o migraciones','Puede retrasar integración y desalinear esquema y API.','Una tarea por rama y PR; revisión compartida de migraciones y compilaciones.','Ambos\nCada semana'],
]
for r,vals in zip(d.tables[7].rows[1:],risks):
    for c,t in zip(r.cells,vals):c.text=t
commit=[
['Acordar alcance, pesos y evidencias por módulo','Integrante 1; revisión Manuel','15-10-2026','9 módulos con pesos que suman 100 % y criterios de cierre acordados.'],
['Demostrar acceso Google y sesión por rol','Manuel: integración. Integrante 1: demo.','15-10-2026','Ingreso y logout correctos; ruta restringida denegada; captura o video y resultado de prueba.'],
]
for r,vals in zip(d.tables[8].rows[1:],commit):
    for c,t in zip(r.cells,vals):c.text=t
# Preserve the supplied form and its four main page blocks.
for st in d.styles:
    if st.type==1:
        st.font.color.rgb=RGBColor(0,0,0)
for ti,table in enumerate(d.tables):
    if ti==2:
        table.autofit=False
        widths=[1.7,1.05,.48,.65,.7,2.5]
        for col,w in zip(table.columns,widths): col.width=Inches(w)
        for row in table.rows:
            for cell,w in zip(row.cells,widths): cell.width=Inches(w)
    for ri,row in enumerate(table.rows):
        pr=row._tr.get_or_add_trPr()
        for el in list(pr):
            if el.tag==qn('w:trHeight'):pr.remove(el)
        ns=OxmlElement('w:cantSplit');pr.append(ns)
        if ri==0:
            rep=OxmlElement('w:tblHeader');pr.append(rep)
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for pp in cell.paragraphs:
                pp.paragraph_format.space_after=Pt(3)
                pp.paragraph_format.space_before=Pt(2)
                for run in pp.runs:
                    run.font.size=Pt(8 if ti==2 else 9)
                    run.font.color.rgb=RGBColor(255,255,255) if ri==0 and ti>0 else RGBColor(0,0,0)
                    run.bold=ri==0
d.core_properties.title='Ficha de avance 04 CRONOS'
d.core_properties.subject='Trazabilidad y planificación de la versión 1.0'
for pp in d.paragraphs:
    if 'Integrante 1' in pp.text: pp.text=pp.text.replace('Integrante 1','Martin Ibarra')
for table in d.tables:
    for row in table.rows:
        for cell in row.cells:
            for pp in cell.paragraphs:
                for run in pp.runs:
                    run.text=run.text.replace('Integrante 1','Martin Ibarra').replace('Manuel','Manuel Araya')
d.save(out/'Ficha_Avance_04_Cronos_COMPLETADA.docx')
print(out/'Ficha_Avance_04_Cronos_COMPLETADA.docx')
